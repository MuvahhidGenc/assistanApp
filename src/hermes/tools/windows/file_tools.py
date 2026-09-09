from __future__ import annotations

import shutil
from pathlib import Path
from typing import Any

from hermes.config.settings import RiskLevel
from hermes.tools.base import BaseTool, ToolExecutionResult
from hermes.tools.windows.input_backend import run_in_thread


def resolve_user_path(raw: str) -> Path:
    """Relative paths anchor to Desktop (e.g. Hermes2/tevhid.docx)."""
    text = (raw or "").strip()
    if not text:
        raise ValueError("path gerekli")
    path = Path(text).expanduser()
    if path.is_absolute():
        return path.resolve()
    from hermes.context.system_paths import current_user_desktop_path

    desktop = current_user_desktop_path()
    if desktop is None:
        raise ValueError("Current user Desktop path could not be observed")
    return (desktop / path).resolve()


def resolve_unique_file_path(path: Path) -> Path:
    """Return path if free, else rapor.txt -> rapor (1).txt -> rapor (2).txt ..."""
    target = path.resolve()
    if not target.exists():
        return target
    parent = target.parent
    stem = target.stem
    suffix = target.suffix
    for index in range(1, 1000):
        candidate = parent / f"{stem} ({index}){suffix}"
        if not candidate.exists():
            return candidate.resolve()
    raise ValueError(f"Benzersiz dosya adi bulunamadi: {target}")


def _dir_listing(folder: Path, limit: int = 50) -> list[dict[str, Any]]:
    if not folder.is_dir():
        return []
    entries: list[dict[str, Any]] = []
    for item in sorted(folder.iterdir(), key=lambda p: p.name.lower())[:limit]:
        entries.append(
            {
                "name": item.name,
                "type": "dir" if item.is_dir() else "file",
                "size": item.stat().st_size if item.is_file() else None,
            }
        )
    return entries


class ReadFileTool(BaseTool):
    name = "read_file"
    description = "Metin dosyasinin icerigini okur."
    risk_level = RiskLevel.READ_ONLY
    category = "files"

    async def execute(self, path: str = "", **kwargs: Any) -> ToolExecutionResult:
        raw = (path or kwargs.get("file") or "").strip()
        if not raw:
            return ToolExecutionResult(success=False, error="path gerekli")
        try:
            target = resolve_user_path(raw)
            if not target.is_file():
                return ToolExecutionResult(success=False, error=f"Dosya bulunamadi: {target}")
            content = target.read_text(encoding="utf-8")
            stat = target.stat()
            return ToolExecutionResult(
                success=True,
                output={
                    "path": str(target),
                    "content": content,
                    "size": stat.st_size,
                    "exists": True,
                },
                verified=True,
            )
        except OSError as exc:
            return ToolExecutionResult(success=False, error=str(exc))

    def get_parameters_schema(self) -> dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "Okunacak dosya yolu"},
            },
            "required": ["path"],
        }


class ListDirectoryTool(BaseTool):
    name = "list_directory"
    description = "Klasor icerigini listeler (dogrulama icin)."
    risk_level = RiskLevel.READ_ONLY
    category = "files"

    async def execute(self, path: str = "", **kwargs: Any) -> ToolExecutionResult:
        raw = (path or kwargs.get("dir") or ".").strip()
        try:
            if raw not in (".", ""):
                target = resolve_user_path(raw)
            else:
                from hermes.context.system_paths import current_user_desktop_path

                target = current_user_desktop_path()
                if target is None:
                    return ToolExecutionResult(
                        success=False,
                        error="Current user Desktop path could not be observed",
                    )
            if not target.exists():
                return ToolExecutionResult(success=False, error=f"Yol yok: {target}")
            if target.is_file():
                stat = target.stat()
                return ToolExecutionResult(
                    success=True,
                    output={
                        "path": str(target),
                        "type": "file",
                        "size": stat.st_size,
                        "exists": True,
                    },
                    verified=True,
                )
            entries = _dir_listing(target)
            return ToolExecutionResult(
                success=True,
                output={"path": str(target), "type": "dir", "entries": entries, "count": len(entries)},
                verified=True,
            )
        except Exception as exc:
            return ToolExecutionResult(success=False, error=str(exc))

    def get_parameters_schema(self) -> dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "Klasor yolu veya Desktop alti goreli yol"},
            },
        }


class WriteFileTool(BaseTool):
    name = "write_file"
    description = "Metin dosyasi olusturur veya uzerine yazar (.txt, .md, .json vb.)."
    risk_level = RiskLevel.NORMAL_MODIFICATION
    category = "files"

    async def execute(
        self,
        path: str = "",
        content: str = "",
        append: bool = False,
        encoding: str = "utf-8",
        **kwargs: Any,
    ) -> ToolExecutionResult:
        raw_path = (path or kwargs.get("file") or "").strip()
        text = content if content else str(kwargs.get("text") or "")
        if not raw_path:
            return ToolExecutionResult(success=False, error="path gerekli")
        try:
            target = resolve_user_path(raw_path)
            suffix = target.suffix.casefold()
            if suffix == ".docx":
                word_tool = CreateWordDocumentTool()
                return await word_tool.execute(
                    path=str(target),
                    content=text,
                    title=str(kwargs.get("title") or ""),
                    **{key: value for key, value in kwargs.items() if key != "title"},
                )
            if suffix == ".pdf":
                return await run_in_thread(_write_real_pdf, target, text)
            from hermes.tools.windows.input_backend import wants_modify_existing, wants_recreate

            user_message = str(kwargs.get("user_message") or "")
            unique_if_exists = kwargs.get("unique_if_exists")
            if unique_if_exists is None:
                unique_if_exists = not wants_modify_existing(user_message)
            else:
                unique_if_exists = bool(unique_if_exists)
            if not append and target.is_file() and not wants_recreate(user_message):
                try:
                    existing = target.read_text(encoding=encoding)
                except OSError:
                    existing = None
                if existing == text:
                    stat = target.stat()
                    return ToolExecutionResult(
                        success=True,
                        output={
                            "path": str(target),
                            "size": stat.st_size,
                            "exists": True,
                            "unchanged": True,
                            "verified_listing": _dir_listing(target.parent, limit=20),
                        },
                        verified=True,
                    )
                if unique_if_exists:
                    target = resolve_unique_file_path(target)
            target.parent.mkdir(parents=True, exist_ok=True)
            if append:
                with target.open("a", encoding=encoding, newline="\n") as handle:
                    handle.write(text)
            else:
                target.write_text(text, encoding=encoding, newline="\n")
            stat = target.stat()
            return ToolExecutionResult(
                success=True,
                output={
                    "path": str(target),
                    "size": stat.st_size,
                    "exists": target.is_file(),
                    "verified_listing": _dir_listing(target.parent, limit=20),
                },
                verified=target.is_file() and stat.st_size >= 0,
            )
        except Exception as exc:
            return ToolExecutionResult(success=False, error=str(exc))

    def get_parameters_schema(self) -> dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "Dosya yolu (or. Hermes2/notlar.txt)"},
                "content": {"type": "string", "description": "Dosya icerigi"},
                "append": {"type": "boolean", "default": False},
            },
            "required": ["path", "content"],
        }


def _pdf_escape(text: str) -> str:
    translit = {
        "ı": "i",
        "İ": "I",
        "ğ": "g",
        "Ğ": "G",
        "ş": "s",
        "Ş": "S",
        "ç": "c",
        "Ç": "C",
        "ö": "o",
        "Ö": "O",
        "ü": "u",
        "Ü": "U",
    }
    cleaned = "".join(translit.get(char, char) for char in text)
    ascii_text = cleaned.encode("latin-1", errors="replace").decode("latin-1")
    return ascii_text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _write_real_pdf(target: Path, text: str) -> ToolExecutionResult:
    from hermes.mission.reality_verification import verify_pdf_document

    target.parent.mkdir(parents=True, exist_ok=True)
    body = (text or "").strip() or "Belge"
    lines = []
    for paragraph in body.replace("\r\n", "\n").split("\n"):
        chunk = paragraph.strip() or " "
        while chunk:
            lines.append(chunk[:90])
            chunk = chunk[90:]
    if not lines:
        lines = ["Belge"]
    commands = ["BT", "/F1 12 Tf", "50 780 Td"]
    for index, line in enumerate(lines[:60]):
        if index:
            commands.append("0 -16 Td")
        commands.append(f"({_pdf_escape(line)}) Tj")
    commands.append("ET")
    stream = "\n".join(commands).encode("latin-1", errors="replace")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, payload in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode("ascii"))
        output.extend(payload)
        output.extend(b"\nendobj\n")
    xref_at = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_at}\n%%EOF\n"
        ).encode("ascii")
    )
    target.write_bytes(bytes(output))
    check = verify_pdf_document(target, expected_text="")
    if not check.get("ok"):
        return ToolExecutionResult(
            success=False,
            error=f"Gecerli PDF yazilamadi: {check.get('reason')}",
            output={"path": str(target), "verified": False, **check},
            verified=False,
        )
    stat = target.stat()
    return ToolExecutionResult(
        success=True,
        output={
            "path": str(target),
            "size": stat.st_size,
            "exists": True,
            "pages": check.get("pages"),
            "format": "pdf",
            "verified_listing": _dir_listing(target.parent, limit=20),
        },
        verified=True,
    )


class CreateWordDocumentTool(BaseTool):
    name = "create_word_document"
    description = "Word (.docx) dosyasi olusturur ve paragraf/metin ekler."
    risk_level = RiskLevel.NORMAL_MODIFICATION
    category = "files"

    async def execute(
        self,
        path: str = "",
        content: str = "",
        title: str = "",
        **kwargs: Any,
    ) -> ToolExecutionResult:
        raw_path = (path or kwargs.get("file") or "").strip()
        if not raw_path.lower().endswith(".docx"):
            raw_path = f"{raw_path}.docx" if raw_path else "document.docx"
        body = content or str(kwargs.get("text") or "")
        heading = (title or kwargs.get("heading") or "").strip()

        def _build() -> dict[str, Any]:
            from docx import Document

            target = resolve_user_path(raw_path)
            target.parent.mkdir(parents=True, exist_ok=True)
            doc = Document()
            if heading:
                doc.add_heading(heading, level=0)
            paragraphs = [part.strip() for part in body.replace("\r\n", "\n").split("\n\n") if part.strip()]
            if not paragraphs and body.strip():
                paragraphs = [line.strip() for line in body.split("\n") if line.strip()]
            for para in paragraphs:
                doc.add_paragraph(para)
            if not paragraphs and not heading:
                doc.add_paragraph("")
            doc.save(str(target))
            from hermes.mission.reality_verification import verify_docx_document

            check = verify_docx_document(target)
            stat = target.stat()
            return {
                "path": str(target),
                "size": stat.st_size,
                "exists": target.is_file(),
                "paragraph_count": len(paragraphs),
                "format_ok": bool(check.get("ok")),
                "format_reason": check.get("reason"),
                "verified_listing": _dir_listing(target.parent, limit=20),
            }

        try:
            data = await run_in_thread(_build)
            format_ok = bool(data.get("format_ok") and data.get("exists"))
            return ToolExecutionResult(
                success=format_ok,
                output=data,
                error="" if format_ok else f"Gecerli DOCX yazilamadi: {data.get('format_reason')}",
                verified=format_ok,
            )
        except ImportError:
            return ToolExecutionResult(
                success=False,
                error="python-docx kurulu degil. pip install python-docx veya client yeniden build edin.",
            )
        except Exception as exc:
            return ToolExecutionResult(success=False, error=str(exc))

    def get_parameters_schema(self) -> dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "Docx yolu (or. Hermes2/tevhid.docx)"},
                "title": {"type": "string", "description": "Baslik (opsiyonel)"},
                "content": {"type": "string", "description": "Govde metni; paragraflar icin bos satir birak"},
            },
            "required": ["path", "content"],
        }


class CopyFileTool(BaseTool):
    name = "copy_file"
    description = "Dosyayi veya klasoru hedefe kopyalar (arka planda, UI acmaz)."
    risk_level = RiskLevel.NORMAL_MODIFICATION
    category = "files"

    async def execute(
        self,
        source: str = "",
        destination: str = "",
        **kwargs: Any,
    ) -> ToolExecutionResult:
        src_raw = (source or kwargs.get("path") or kwargs.get("from") or "").strip()
        dst_raw = (destination or kwargs.get("dest") or kwargs.get("to") or "").strip()
        if not src_raw or not dst_raw:
            return ToolExecutionResult(success=False, error="source ve destination gerekli")
        try:
            src = resolve_user_path(src_raw)
            dst = resolve_user_path(dst_raw)
            if not src.exists():
                return ToolExecutionResult(success=False, error=f"Kaynak yok: {src}")
            if (
                dst.is_dir()
                or str(dst_raw).endswith(("/", "\\"))
                or (not dst.exists() and not dst.suffix)
            ):
                dst = dst / src.name
            dst.parent.mkdir(parents=True, exist_ok=True)
            if src.is_dir():
                if dst.exists():
                    shutil.copytree(src, dst, dirs_exist_ok=True)
                else:
                    shutil.copytree(src, dst)
            else:
                shutil.copy2(src, dst)
            exists = dst.exists()
            return ToolExecutionResult(
                success=exists,
                output={
                    "source": str(src),
                    "destination": str(dst),
                    "exists": exists,
                    "verified_listing": _dir_listing(dst.parent, limit=20),
                },
                verified=exists,
            )
        except Exception as exc:
            return ToolExecutionResult(success=False, error=str(exc))

    def get_parameters_schema(self) -> dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "source": {"type": "string"},
                "destination": {"type": "string"},
            },
            "required": ["source", "destination"],
        }


class MoveFileTool(BaseTool):
    name = "move_file"
    description = "Dosyayi veya klasoru tasir (arka planda, UI acmaz)."
    risk_level = RiskLevel.NORMAL_MODIFICATION
    category = "files"

    async def execute(
        self,
        source: str = "",
        destination: str = "",
        **kwargs: Any,
    ) -> ToolExecutionResult:
        src_raw = (source or kwargs.get("path") or kwargs.get("from") or "").strip()
        dst_raw = (destination or kwargs.get("dest") or kwargs.get("to") or "").strip()
        if not src_raw or not dst_raw:
            return ToolExecutionResult(success=False, error="source ve destination gerekli")
        try:
            src = resolve_user_path(src_raw)
            dst = resolve_user_path(dst_raw)
            if not src.exists():
                return ToolExecutionResult(success=False, error=f"Kaynak yok: {src}")
            if dst.is_dir() or str(dst_raw).endswith(("/", "\\")):
                dst = dst / src.name
            dst.parent.mkdir(parents=True, exist_ok=True)
            shutil.move(str(src), str(dst))
            src_gone = not src.exists()
            dst_exists = dst.exists()
            return ToolExecutionResult(
                success=src_gone and dst_exists,
                output={
                    "source": str(src),
                    "destination": str(dst),
                    "source_exists": src.exists(),
                    "destination_exists": dst_exists,
                    "verified_listing": _dir_listing(dst.parent, limit=20),
                },
                verified=src_gone and dst_exists,
            )
        except Exception as exc:
            return ToolExecutionResult(success=False, error=str(exc))

    def get_parameters_schema(self) -> dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "source": {"type": "string"},
                "destination": {"type": "string"},
            },
            "required": ["source", "destination"],
        }


class RenamePathTool(BaseTool):
    name = "rename_path"
    description = "Dosya veya klasor adini degistirir."
    risk_level = RiskLevel.NORMAL_MODIFICATION
    category = "files"

    async def execute(
        self,
        path: str = "",
        new_name: str = "",
        **kwargs: Any,
    ) -> ToolExecutionResult:
        raw = (path or kwargs.get("source") or "").strip()
        name = (new_name or kwargs.get("name") or kwargs.get("to") or "").strip()
        if not raw or not name:
            return ToolExecutionResult(success=False, error="path ve new_name gerekli")
        try:
            src = resolve_user_path(raw)
            if not src.exists():
                return ToolExecutionResult(success=False, error=f"Yol yok: {src}")
            dst = src.with_name(name)
            src.rename(dst)
            return ToolExecutionResult(
                success=dst.exists() and not src.exists(),
                output={
                    "source": str(src),
                    "destination": str(dst),
                    "exists": dst.exists(),
                    "verified_listing": _dir_listing(dst.parent, limit=20),
                },
                verified=dst.exists(),
            )
        except Exception as exc:
            return ToolExecutionResult(success=False, error=str(exc))

    def get_parameters_schema(self) -> dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "path": {"type": "string"},
                "new_name": {"type": "string"},
            },
            "required": ["path", "new_name"],
        }


class SearchFilesTool(BaseTool):
    name = "search_files"
    description = "Klasorde dosya arar (glob ve opsiyonel tarih filtresi)."
    risk_level = RiskLevel.READ_ONLY
    category = "files"

    async def execute(
        self,
        path: str = "",
        pattern: str = "*",
        modified_within_hours: int | None = None,
        **kwargs: Any,
    ) -> ToolExecutionResult:
        raw = (path or kwargs.get("folder") or kwargs.get("directory") or "").strip()
        glob_pat = (pattern or kwargs.get("glob") or "*").strip()
        try:
            folder = resolve_user_path(raw) if raw else Path.home() / "Downloads"
            if not folder.is_dir():
                return ToolExecutionResult(success=False, error=f"Klasor yok: {folder}")
            from datetime import datetime, timedelta, timezone

            cutoff = None
            hours = modified_within_hours or kwargs.get("hours")
            if hours is not None:
                cutoff = datetime.now(timezone.utc) - timedelta(hours=int(hours))
            matches: list[dict[str, Any]] = []
            for item in folder.rglob(glob_pat.lstrip("/")):
                if not item.is_file():
                    continue
                if cutoff is not None:
                    mtime = datetime.fromtimestamp(item.stat().st_mtime, tz=timezone.utc)
                    if mtime < cutoff:
                        continue
                matches.append(
                    {
                        "path": str(item),
                        "name": item.name,
                        "size": item.stat().st_size,
                        "modified": item.stat().st_mtime,
                    }
                )
            return ToolExecutionResult(
                success=True,
                output={
                    "folder": str(folder),
                    "pattern": glob_pat,
                    "count": len(matches),
                    "matches": matches[:100],
                },
                verified=True,
            )
        except Exception as exc:
            return ToolExecutionResult(success=False, error=str(exc))

    def get_parameters_schema(self) -> dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "path": {"type": "string"},
                "pattern": {"type": "string", "default": "*"},
                "modified_within_hours": {"type": "integer"},
            },
        }


class DeletePathTool(BaseTool):
    name = "delete_path"
    description = "Dosya veya klasoru siler (klasor icin recursive=true)."
    risk_level = RiskLevel.HIGH_RISK
    category = "files"

    async def execute(
        self,
        path: str = "",
        recursive: bool = False,
        **kwargs: Any,
    ) -> ToolExecutionResult:
        raw = (path or kwargs.get("target") or "").strip()
        if not raw:
            return ToolExecutionResult(success=False, error="path gerekli")
        try:
            from hermes.mission.reality_verification import (
                find_office_lock_files,
                is_office_lock_name,
                office_lock_stem_hint,
            )

            target = resolve_user_path(raw)
            parent = target.parent
            name = target.name
            lock_targets: list[Path] = []
            if is_office_lock_name(name) or not target.exists():
                stem = office_lock_stem_hint(name)
                lock_targets = find_office_lock_files(parent, stem)
                if target.exists() and is_office_lock_name(name):
                    lock_targets = [target, *[item for item in lock_targets if item != target]]
            elif target.exists() and target.is_file() and not is_office_lock_name(name):
                siblings = find_office_lock_files(parent, office_lock_stem_hint(target.name))
                if siblings:
                    lock_targets = siblings
            if lock_targets:
                deleted: list[str] = []
                remaining: list[str] = []
                errors: list[str] = []
                for item in lock_targets:
                    try:
                        item.unlink()
                    except OSError as exc:
                        errors.append(f"{item.name}: {exc}")
                    if item.exists():
                        remaining.append(item.name)
                    else:
                        deleted.append(item.name)
                cleaned = bool(deleted) and not remaining
                return ToolExecutionResult(
                    success=cleaned,
                    error="" if cleaned else (
                        "Kilit dosyalari temizlenemedi: " + "; ".join(errors + remaining)
                        if remaining or errors
                        else "Kilit dosyasi bulunamadi"
                    ),
                    output={
                        "path": str(parent),
                        "deleted": deleted,
                        "remaining": remaining,
                        "exists_after": bool(remaining),
                        "verified_listing": _dir_listing(parent, limit=20) if parent.is_dir() else [],
                    },
                    verified=cleaned,
                )
            if not target.exists():
                return ToolExecutionResult(success=False, error=f"Yol bulunamadi: {target}")
            if target.is_dir():
                if recursive:
                    shutil.rmtree(target)
                else:
                    target.rmdir()
            else:
                target.unlink()
            still_exists = target.exists()
            return ToolExecutionResult(
                success=not still_exists,
                output={
                    "deleted": name,
                    "path": str(target),
                    "exists_after": still_exists,
                    "verified_listing": _dir_listing(parent, limit=20) if parent.is_dir() else [],
                },
                verified=not still_exists,
            )
        except Exception as exc:
            return ToolExecutionResult(success=False, error=str(exc))

    def get_parameters_schema(self) -> dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "Silinecek dosya/klasor yolu"},
                "recursive": {
                    "type": "boolean",
                    "default": False,
                    "description": "Dolu klasor silmek icin true",
                },
            },
            "required": ["path"],
        }
